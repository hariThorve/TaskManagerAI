import streamlit as st
from langchain_groq import ChatGroq
from langchain_core.prompts import ChatPromptTemplate
from datetime import datetime, timedelta
import time
import os
from dotenv import load_dotenv
import json
import pandas as pd
import sqlite3
import uuid
import plotly.express as px
import io
from openpyxl import Workbook

# Load environment variables from .env file
load_dotenv()

# Set up groq api key
api = os.getenv("GROQ_API_KEY")

# Initialize ChatGroq model
chatgroq_model = ChatGroq(api_key=api)

# Create prompt templates
completion_prompt_template = ChatPromptTemplate.from_messages([
    ("system", "You are an aggressively motivating assistant in the style of David Goggins. Generate a powerful, intense congratulatory message for completing a task on time. Use strong language but maintain a positive tone. Keep the response to 2-3 impactful sentences."),
    ("user", "Task '{task_name}' completed on time! Generate a motivational congratulatory response.")
])

wakeup_prompt_template = ChatPromptTemplate.from_messages([
    ("system", "You are an aggressively motivating assistant in the style of David Goggins. Generate a wake-up call message for tasks that are either incomplete or completed late. Use strong language to push them to do better. Keep the response to 2-3 intense sentences."),
    ("user", "Task '{task_name}' was {status}. Generate a wake-up call message.")
])

chat_prompt_template = ChatPromptTemplate.from_messages([
    ("system", "You are an aggressively motivating AI assistant inspired by David Goggins. You provide tough love, motivation, and direct answers with no sugar coating. You can use strong language when appropriate but maintain helpfulness. Your responses should be intense yet constructive. Complete your sentences in 2 to 3 lines. If user asks for detailed answer then you can give him detailed answer."),
    ("user", "{user_input}")
])

def apply_premium_styling():
    st.markdown("""
        <style>
        /* Premium Dark Theme with Emerald Accents */
        .stApp {
            background: linear-gradient(135deg, #1a1a2e 0%, #102b2b 50%, #1a1a2e 100%);
            color: #e0e0e0;
            font-family: 'Inter', system-ui, -apple-system, sans-serif;
        }

        /* Enhanced Headers */
        h1, h2, h3 {
            background: linear-gradient(90deg, #20c997, #155e63);
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
            font-weight: 800;
            letter-spacing: -0.5px;
            animation: gradientFlow 3s ease infinite;
        }

        /* Glassmorphism Cards */
        .element-container, div.stButton > button, .stTextInput > div {
            background: rgba(255, 255, 255, 0.05);
            backdrop-filter: blur(10px);
            border: 1px solid rgba(255, 255, 255, 0.1);
            border-radius: 10px;
            transition: all 0.3s ease;
        }

        /* Enhanced Buttons */
        div.stButton > button {
            background: linear-gradient(45deg, #20c997, #155e63);
            color: white;
            font-weight: 600;
            border: none;
            padding: 0.5rem 1rem;
            border-radius: 8px;
            transition: all 0.3s ease;
            text-transform: uppercase;
            letter-spacing: 1px;
        }

        div.stButton > button:hover {
            transform: translateY(-2px);
            box-shadow: 0 5px 15px rgba(32, 201, 151, 0.4);
        }

        /* Input Fields */
        .stTextInput > div > div {
            background: rgba(255, 255, 255, 0.05);
            border: 1px solid rgba(255, 255, 255, 0.1);
            border-radius: 8px;
            color: white;
        }

        .stTextInput > div > div:focus-within {
            border-color: #20c997;
            box-shadow: 0 0 10px rgba(32, 201, 151, 0.3);
            transform: translateY(-1px);
        }

        /* Task Cards */
        .task-card {
            background: rgba(255, 255, 255, 0.05);
            border-radius: 10px;
            padding: 1rem;
            margin: 1rem 0;
            border: 1px solid rgba(255, 255, 255, 0.1);
            transition: all 0.3s ease;
            animation: fadeIn 0.5s ease-out;
        }

        .task-card:hover {
            transform: translateY(-2px);
            border-color: #20c997;
            box-shadow: 0 5px 15px rgba(32, 201, 151, 0.2);
        }

        /* Other styles remain the same but using emerald and dark teal colors */
        </style>
    """, unsafe_allow_html=True)

class DatabaseManager:
    def __init__(self):
        self.init_database()

    def init_database(self):
        conn = self.get_connection()
        c = conn.cursor()
        
         # Create tasks table
        c.execute('''
            CREATE TABLE IF NOT EXISTS tasks (
                id TEXT PRIMARY KEY,
                task TEXT NOT NULL,
                time DATETIME NOT NULL,
                status TEXT NOT NULL,
                priority TEXT NOT NULL,
                category TEXT NOT NULL,
                notes TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (category) REFERENCES categories(name)
            )
        ''')
        
        # Create categories table
        c.execute('''
            CREATE TABLE IF NOT EXISTS categories (
                id TEXT PRIMARY KEY,
                name TEXT UNIQUE NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        
        # Create chat_history table
        c.execute('''
            CREATE TABLE IF NOT EXISTS chat_history (
                id TEXT PRIMARY KEY,
                role TEXT NOT NULL,
                content TEXT NOT NULL,
                timestamp TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        
        conn.commit()
        conn.close()

    def ensure_default_category(self):
        conn = self.get_connection()
        c = conn.cursor()
        
        # Check if any categories exist
        c.execute("SELECT COUNT(*) FROM categories")
        category_count = c.fetchone()[0]
        
        # If no categories exist, add default category
        if category_count == 0:
            default_category_id = str(uuid.uuid4())
            c.execute(
                "INSERT INTO categories (id, name) VALUES (?, ?)",
                (default_category_id, "General")
            )
            conn.commit()
        
        conn.close()    

    @staticmethod
    def get_connection():
        return sqlite3.connect('goggins_bot.db')

    def save_task(self, task):
        # Validate category exists
        conn = self.get_connection()
        c = conn.cursor()
        
        c.execute("SELECT name FROM categories WHERE name = ?", (task['category'],))
        category_exists = c.fetchone()
        
        if not category_exists:
            conn.close()
            raise ValueError(f"Category '{task['category']}' does not exist!")
        
        task_id = str(uuid.uuid4())
        try:
            c.execute('''
                INSERT INTO tasks (id, task, time, status, priority, category, notes)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            ''', (task_id, task['task'], task['time'], task['status'], 
                  task['priority'], task['category'], task['notes']))
            conn.commit()
        except sqlite3.IntegrityError as e:
            conn.close()
            raise ValueError(f"Error saving task: {str(e)}")
        
        conn.close()
        return task_id

    def get_tasks(self, filter_completed=False, filter_category=None, filter_priority=None):
        conn = self.get_connection()
        query = "SELECT * FROM tasks WHERE 1=1"
        params = []
        
        if not filter_completed:
            query += " AND status != 'completed'"
        
        if filter_category:
            placeholders = ','.join(['?' for _ in filter_category])
            query += f" AND category IN ({placeholders})"
            params.extend(filter_category)
            
        if filter_priority:
            placeholders = ','.join(['?' for _ in filter_priority])
            query += f" AND priority IN ({placeholders})"
            params.extend(filter_priority)
            
        df = pd.read_sql_query(query, conn, params=params)
        conn.close()
        return df

    def update_task_status(self, task_id, status):
        conn = self.get_connection()
        c = conn.cursor()
        
        # Get task details before updating
        c.execute("SELECT task, time FROM tasks WHERE id = ?", (task_id,))
        task_data = c.fetchone()
        task_name = task_data[0]
        task_time = datetime.strptime(task_data[1], "%Y-%m-%d %H:%M")
        
        # Update task status
        c.execute("UPDATE tasks SET status = ? WHERE id = ?", (status, task_id))
        conn.commit()
        conn.close()

        current_time = datetime.now()
        
        # Generate appropriate message based on completion time
        if status == 'completed':
            if current_time <= task_time:
                # Task completed on time
                prompt = completion_prompt_template.format_messages(task_name=task_name)
                response = chatgroq_model(prompt)
                return response.content, 'success'
            else:
                # Task completed late
                prompt = wakeup_prompt_template.format_messages(
                    task_name=task_name,
                    status='completed late'
                )
                response = chatgroq_model(prompt)
                return response.content, 'warning'
        else:
            # Task marked as incomplete
            prompt = wakeup_prompt_template.format_messages(
                task_name=task_name,
                status='not completed'
            )
            response = chatgroq_model(prompt)
            return response.content, 'warning'

    def save_category(self, category_name):
        conn = self.get_connection()
        c = conn.cursor()
        category_id = str(uuid.uuid4())
        try:
            c.execute("INSERT INTO categories (id, name) VALUES (?, ?)", 
                     (category_id, category_name))
            conn.commit()
        except sqlite3.IntegrityError:
            pass  # Category already exists
        conn.close()

    def get_categories(self):
        conn = self.get_connection()
        c = conn.cursor()
        c.execute("SELECT name FROM categories")
        categories = [row[0] for row in c.fetchall()]
        conn.close()
        return categories

    def get_analytics_data(self):
        conn = self.get_connection()
        
        # Get all tasks
        tasks_df = pd.read_sql_query("SELECT * FROM tasks", conn)
        tasks_df['time'] = pd.to_datetime(tasks_df['time'])
        tasks_df['created_at'] = pd.to_datetime(tasks_df['created_at'])
        
        # Calculate completion rates by category
        category_completion = tasks_df.groupby('category').agg({
            'status': lambda x: (x == 'completed').mean() * 100
        }).round(2)
        
        # Calculate daily task counts
        daily_tasks = tasks_df.groupby(tasks_df['created_at'].dt.date).size().reset_index()
        daily_tasks.columns = ['date', 'count']
        
        # Get recent tasks
        now = datetime.now()
        recent_tasks = tasks_df[tasks_df['time'] >= (now - timedelta(days=7))]
        previous_week_tasks = tasks_df[
            (tasks_df['time'] >= (now - timedelta(days=14))) &
            (tasks_df['time'] < (now - timedelta(days=7)))
        ]
        
        conn.close()
        
        return {
            'tasks_df': tasks_df,
            'category_completion': category_completion,
            'daily_tasks': daily_tasks,
            'recent_tasks': recent_tasks,
            'previous_week_tasks': previous_week_tasks
        }

    @staticmethod
    def save_chat_message(role, content):
        conn = DatabaseManager.get_connection()
        conn.execute("INSERT INTO chat_history (id, role, content) VALUES (?, ?, ?)",
                     (str(uuid.uuid4()), role, content))
        conn.commit()
        conn.close()

    @staticmethod
    def get_chat_history():
        conn = DatabaseManager.get_connection()
        df = pd.read_sql_query("SELECT * FROM chat_history ORDER BY timestamp ASC", conn)
        conn.close()
        return df.to_dict('records')

    def clear_chat_history(self):
        conn = self.get_connection()
        c = conn.cursor()
        c.execute('DELETE FROM chat_history')
        conn.commit()
        conn.close()

def init_session_state():
    if 'db' not in st.session_state:
        st.session_state.db = DatabaseManager()
    if 'last_response' not in st.session_state:
        st.session_state.last_response = None
    if 'response_type' not in st.session_state:
        st.session_state.response_type = None

def show_task_manager():
    st.title("📋 Task Manager - STAY HARD!")
    
    # Add Category Section
    with st.expander("Manage Categories 📑"):
        with st.form(key='category_form'):
            new_category = st.text_input("Add New Category:")
            add_category = st.form_submit_button("ADD CATEGORY 💪")
            if add_category and new_category:
                try:
                    categories = st.session_state.db.get_categories()
                    if new_category not in categories:
                        st.session_state.db.save_category(new_category)
                        st.success(f"NEW CATEGORY ADDED! LET'S GO! 🔥")
                        st.rerun()
                    else:
                        st.warning("Category already exists!")
                except Exception as e:
                    st.error(f"Error adding category: {str(e)}")
    
    # Get categories before task form
    categories = st.session_state.db.get_categories()
    
    # Show warning if no categories exist
    if not categories:
        st.error("NO CATEGORIES FOUND! ADD A CATEGORY FIRST! 💪")
        return
    
    # Task input form
    with st.form(key='task_form'):
        st.subheader("Add New Task")
        task_name = st.text_input("Task Name:")
        col_a, col_b = st.columns(2)
        
        with col_a:
            task_date = st.date_input("Date:")
            task_priority = st.select_slider("Priority", options=['Low', 'Medium', 'High'], value='Medium')
        
        with col_b:
            task_time = st.time_input("Time:")
            task_category = st.selectbox("Category", categories)
        
        task_notes = st.text_area("Notes (optional):")
        submit_button = st.form_submit_button(label='SET THIS TASK! 💪')
        
        if submit_button:
            if not task_name:
                st.error("TASK NAME IS REQUIRED! DON'T BE SOFT! 💪")
            else:
                try:
                    task_datetime = datetime.combine(task_date, task_time)
                    new_task = {
                        "task": task_name,
                        "time": task_datetime.strftime("%Y-%m-%d %H:%M"),
                        "status": "pending",
                        "priority": task_priority,
                        "category": task_category,
                        "notes": task_notes
                    }
                    st.session_state.db.save_task(new_task)
                    st.success(f"TASK SET! NO EXCUSES NOW! 🔥")
                    st.rerun()
                except Exception as e:
                    st.error(f"Error saving task: {str(e)}")

    # Task filtering
    st.subheader("Filter Tasks")
    col1, col2 = st.columns(2)
    with col1:
        filter_category = st.multiselect("Filter by Category", categories)
    with col2:
        filter_priority = st.multiselect("Filter by Priority", ['Low', 'Medium', 'High'])
    show_completed = st.checkbox("Show completed tasks")

    # Display tasks
    st.subheader("YOUR BATTLE PLAN:")
    try:
        tasks_df = st.session_state.db.get_tasks(
            filter_completed=show_completed,
            filter_category=filter_category if filter_category else None,
            filter_priority=filter_priority if filter_priority else None
        )
        
        if tasks_df.empty:
            st.info("NO TASKS FOUND WITH CURRENT FILTERS! TIME TO ADD SOME! 💪")
            return
        
        # Sort tasks by date and priority
        tasks_df['time'] = pd.to_datetime(tasks_df['time'])
        tasks_df = tasks_df.sort_values(['time', 'priority'], 
                                      ascending=[True, False])
        
        for _, task in tasks_df.iterrows():
            with st.container():
                col1, col2, col3, col4 = st.columns([3, 2, 1, 1])
                
                with col1:
                    # Task name and priority indicator
                    task_header = f"**{task['task']}**"
                    if task['priority'] == 'High':
                        task_header += " 🔥"
                    elif task['priority'] == 'Medium':
                        task_header += " ⚡"
                    st.markdown(task_header)
                    
                    # Category and notes
                    st.caption(f"Category: {task['category']}")
                    if task['notes']:
                        with st.expander("Notes"):
                            st.write(task['notes'])
                
                with col2:
                    # Format datetime for display
                    task_time = pd.to_datetime(task['time'])
                    st.write(task_time.strftime("%Y-%m-%d %H:%M"))
                    
                    # Show status indicator
                    if task['status'] == 'completed':
                        st.success("Completed ✓")
                    elif task_time < pd.Timestamp.now():
                        st.error("Overdue!")
                    else:
                        st.info("Pending")
                
                with col3:
                    st.write(f"Priority: {task['priority']}")
                
                with col4:
                    if task['status'] != 'completed':
                        if st.button("Complete ✓", key=f"complete_{task['id']}"):
                            try:
                                message, message_type = st.session_state.db.update_task_status(
                                    task['id'], 
                                    'completed'
                                )
                                st.session_state.last_response = message
                                st.session_state.response_type = message_type
                                st.rerun()
                            except Exception as e:
                                st.error(f"Error updating task: {str(e)}")
                
                st.divider()
                
    except Exception as e:
        st.error(f"Error loading tasks: {str(e)}")

    # Add export functionality
    if not tasks_df.empty:
        st.subheader("Export Tasks")
        export_format = st.selectbox(
            "Select export format:",
            ["Excel", "CSV", "JSON"],
            key="export_format"
        )
        
        if st.button("EXPORT TASKS 📊"):
            try:
                if export_format == "Excel":
                    buffer = io.BytesIO()
                    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
                        tasks_df.to_excel(writer, index=False)
                    st.download_button(
                        label="Download Excel",
                        data=buffer.getvalue(),
                        file_name="tasks_export.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
                elif export_format == "CSV":
                    csv = tasks_df.to_csv(index=False)
                    st.download_button(
                        label="Download CSV",
                        data=csv,
                        file_name="tasks_export.csv",
                        mime="text/csv"
                    )
                else:  # JSON
                    json_str = tasks_df.to_json(orient='records', date_format='iso')
                    st.download_button(
                        label="Download JSON",
                        data=json_str,
                        file_name="tasks_export.json",
                        mime="application/json"
                    )
                st.success("EXPORT READY! GET AFTER IT! 💪")
            except Exception as e:
                st.error(f"Error exporting tasks: {str(e)}")

def show_analytics():
    st.title("📊 Task Analytics - TRACK YOUR PROGRESS!")
    
    try:
        # Get analytics data
        analytics_data = st.session_state.db.get_analytics_data()
        
        if analytics_data['tasks_df'].empty:
            st.warning("NO DATA TO ANALYZE YET! START ADDING TASKS, WARRIOR! 💪")
            return
        
        tab1, tab2, tab3 = st.tabs(["Overview", "Detailed Analysis", "Time Trends"])
        
        with tab1:
            st.subheader("Overall Performance")
            
            # Calculate key metrics
            df = analytics_data['tasks_df']
            total_tasks = len(df)
            completed_tasks = len(df[df['status'] == 'completed'])
            overdue_tasks = len(df[
                (df['status'] != 'completed') & 
                (pd.to_datetime(df['time']) < pd.Timestamp.now())
            ])
            pending_tasks = total_tasks - completed_tasks
            completion_rate = (completed_tasks / total_tasks * 100) if total_tasks > 0 else 0
            
            # Display metrics
            col1, col2, col3, col4 = st.columns(4)
            with col1:
                st.metric("Total Tasks", total_tasks)
            with col2:
                st.metric("Completed", completed_tasks)
            with col3:
                st.metric("Pending", pending_tasks)
            with col4:
                st.metric("Overdue", overdue_tasks)
            
            # Display completion rate with motivation
            st.metric("Overall Completion Rate", f"{completion_rate:.1f}%")
            if completion_rate >= 80:
                st.success("CRUSHING IT! KEEP PUSHING! 🔥")
            elif completion_rate >= 50:
                st.info("GOOD PROGRESS! BUT YOU CAN DO BETTER! 💪")
            else:
                st.warning("TIME TO STEP IT UP! NO EXCUSES! 😤")
        
        with tab2:
            st.subheader("Task Distribution")
            
            col1, col2 = st.columns(2)
            with col1:
                # Category distribution
                category_counts = df['category'].value_counts()
                fig_category = px.pie(
                    values=category_counts.values,
                    names=category_counts.index,
                    title="Tasks by Category",
                    hole=0.4
                )
                fig_category.update_traces(textposition='inside', textinfo='percent+label')
                st.plotly_chart(fig_category)
            
            with col2:
                # Priority distribution
                priority_counts = df['priority'].value_counts()
                fig_priority = px.pie(
                    values=priority_counts.values,
                    names=priority_counts.index,
                    title="Tasks by Priority",
                    hole=0.4,
                    color_discrete_map={'High': 'red', 'Medium': 'orange', 'Low': 'blue'}
                )
                fig_priority.update_traces(textposition='inside', textinfo='percent+label')
                st.plotly_chart(fig_priority)
            
            # Category performance
            st.subheader("Category Performance")
            category_completion = analytics_data['category_completion']
            fig_completion = px.bar(
                category_completion,
                title="Completion Rate by Category",
                labels={'status': 'Completion Rate (%)', 'category': 'Category'},
                color_discrete_sequence=['#00CED1']  # Turquoise color
            )
            fig_completion.update_layout(showlegend=False)
            st.plotly_chart(fig_completion)
        
        with tab3:
            st.subheader("Time Analysis")
            
            # Task creation trend
            daily_tasks = analytics_data['daily_tasks']
            fig_daily = px.line(
                daily_tasks,
                x='date',
                y='count',
                title="Daily Task Creation Trend",
                labels={'count': 'Number of Tasks', 'date': 'Date'}
            )
            fig_daily.update_traces(line_color='#00CED1')
            st.plotly_chart(fig_daily)
            
            # Week comparison
            st.subheader("Week-over-Week Comparison")
            recent_tasks = analytics_data['recent_tasks']
            previous_tasks = analytics_data['previous_week_tasks']
            
            col1, col2 = st.columns(2)
            with col1:
                recent_count = len(recent_tasks)
                previous_count = len(previous_tasks)
                week_change = ((recent_count - previous_count) / previous_count * 100) if previous_count > 0 else 0
                
                st.metric(
                    "Task Volume",
                    f"{recent_count} tasks this week",
                    f"{week_change:+.1f}% vs last week"
                )
            
            with col2:
                recent_completion = (len(recent_tasks[recent_tasks['status'] == 'completed']) / recent_count * 100) if recent_count > 0 else 0
                previous_completion = (len(previous_tasks[previous_tasks['status'] == 'completed']) / previous_count * 100) if previous_count > 0 else 0
                completion_change = recent_completion - previous_completion
                
                st.metric(
                    "Completion Rate",
                    f"{recent_completion:.1f}% this week",
                    f"{completion_change:+.1f}% vs last week"
                )

    except Exception as e:
        st.error(f"Error generating analytics: {str(e)}")


def show_chat():
    st.title("💪 GOGGINS BOT - NO WEAKNESS HERE!")
    
    try:
        # Chat history display
        chat_history = st.session_state.db.get_chat_history()
        
        # Custom CSS for chat interface
        st.markdown("""
            <style>
            .chat-wrapper {
                max-width: 850px;
                margin: 0 auto;
                padding: 20px;
            }
            
            .chat-container {
                background: #2d2d3d;
                border-radius: 12px;
                box-shadow: 0 4px 6px rgba(0, 0, 0, 0.1);
                overflow: hidden;
            }
            
            .chat-header {
                background: #363648;
                padding: 20px;
                border-bottom: 1px solid #3d3d4d;
                display: flex;
                align-items: center;
            }
            
            .status-indicator {
                width: 12px;
                height: 12px;
                background: #4a9eff;
                border-radius: 50%;
                margin-right: 15px;
            }
            
            .chat-title {
                color: white;
                font-size: 1.2rem;
                font-weight: 600;
                margin: 0;
            }
            
            .chat-subtitle {
                color: #9ca3af;
                font-size: 0.875rem;
                margin-top: 4px;
            }
            
            .messages-container {
                padding: 20px;
                height: 10px;
                overflow-y: auto;
            }
            
            .message {
                margin-bottom: 20px;
                display: flex;
                flex-direction: column;
            }
            
            .message-user {
                align-items: flex-end;
            }
            
            .message-bot {
                align-items: flex-start;
            }
            
            .message-content {
                max-width: 70%;
                padding: 12px 16px;
                border-radius: 12px;
                margin-bottom: 4px;
            }
            
            .message-user .message-content {
                background: #4a9eff;
                color: white;
            }
            
            .message-bot .message-content {
                background: #363648;
                color: white;
            }
            
            .message-timestamp {
                font-size: 0.75rem;
                color: #9ca3af;
                margin-top: 4px;
            }
            
            .chat-input-container {
                background: #363648;
                padding: 20px;
                border-top: 1px solid #3d3d4d;
            }
            
            .chat-input-wrapper {
                display: flex;
                gap: 12px;
            }
            
            .chat-input {
                flex: 1;
                background: #2d2d3d;
                border: 1px solid #3d3d4d;
                border-radius: 8px;
                padding: 12px 16px;
                color: white;
            }
            
            .chat-input:focus {
                outline: none;
                border-color: #4a9eff;
                box-shadow: 0 0 0 2px rgba(74, 158, 255, 0.2);
            }
            
            .chat-send-button {
                background: #4a9eff;
                color: white;
                border: none;
                border-radius: 8px;
                padding: 12px 24px;
                font-weight: 600;
                cursor: pointer;
                transition: background-color 0.2s;
            }
            
            .chat-send-button:hover {
                background: #3a8eef;
            }
            
            .clear-chat-button {
                margin-top: 12px;
                background: #363648;
                color: #9ca3af;
                border: 1px solid #3d3d4d;
                border-radius: 8px;
                padding: 8px 16px;
                font-size: 0.875rem;
                cursor: pointer;
                transition: all 0.2s;
            }
            
            .clear-chat-button:hover {
                background: #3d3d4d;
                color: white;
            }
            </style>
            """, unsafe_allow_html=True)
        
        # Chat Interface Structure
        st.markdown("""
            <div class="chat-wrapper">
                <div class="chat-container">
                    <div class="chat-header">
                        <div class="status-indicator"></div>
                        <div>
                            <h2 class="chat-title">Goggins Bot</h2>
                            <p class="chat-subtitle">Your Personal Motivational Assistant</p>
                        </div>
                    </div>
                    <div class="messages-container">
        """, unsafe_allow_html=True)
        
        # Display Messages
        for message in chat_history:
            message_class = "message-user" if message['role'] == 'user' else "message-bot"
            st.markdown(f"""
                <div class="message {message_class}">
                    <div class="message-content">
                        {message['content']}
                    </div>
                    <div class="message-timestamp">
                        {pd.to_datetime(message['timestamp']).strftime('%I:%M %p')}
                    </div>
                </div>
            """, unsafe_allow_html=True)
        
        st.markdown("""
                    </div>
                </div>
            </div>
        """, unsafe_allow_html=True)
        
        # Chat Input Form
        with st.form(key='chat_form', clear_on_submit=True):
            col1, col2 = st.columns([4, 1])
            
            with col1:
                user_input = st.text_input("", 
                    placeholder="Type your message here...",
                    label_visibility="collapsed")
            
            with col2:
                send_message = st.form_submit_button("Send")
            
            clear_chat = st.form_submit_button("Clear Chat")
            
            if send_message and user_input:
                DatabaseManager.save_chat_message('user', user_input)
                prompt = chat_prompt_template.format_messages(user_input=user_input)
                response = chatgroq_model(prompt)
                DatabaseManager.save_chat_message('assistant', response.content)
                st.rerun()
            
            if clear_chat:
                st.session_state.db.clear_chat_history()
                st.rerun()
                
    except Exception as e:
        st.error(f"Error in chat interface: {str(e)}")

def main():
    st.set_page_config(
        page_title="Goggins Task Manager",
        page_icon="💪",
        layout="wide"
    )
    
    init_session_state()
    apply_premium_styling()
    
    # Display any pending responses
    if st.session_state.last_response:
        if st.session_state.response_type == 'success':
            st.success(st.session_state.last_response)
        else:
            st.warning(st.session_state.last_response)
        st.session_state.last_response = None
        st.session_state.response_type = None
    
    # Navigation
    st.sidebar.title("NAVIGATE, WARRIOR! 💪")
    page = st.sidebar.radio("Choose Your Battle:", 
                           ["Task Manager", "Analytics", "Chat with Goggins"])
    
    if page == "Task Manager":
        show_task_manager()
    elif page == "Analytics":
        show_analytics()
    else:
        show_chat()

if __name__ == "__main__":
    main()